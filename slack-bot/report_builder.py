"""
Builds a multi-tab digital competitive report (one tab per race, plus a
Summary index) from an AdHawk CSV export. Generalized from the one-off
PA report so it keeps working as new exports arrive with different
races, different week ranges, and possibly non-partisan spenders.
"""

import re
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU

LOGO_PATH = Path(__file__).parent / 'assets' / 'gps_impact_logo_white.png'
LOGO_HEIGHT_PX = 40  # header bar is ~69px tall (34pt + 18pt); leaves margin above/below
LOGO_RIGHT_MARGIN_PX = 16

PLATFORM_ORDER = {'CTV': 0, 'Google': 1, 'Facebook': 2, 'Twitter': 3}
SUMMARY_PLATFORMS = ['CTV', 'Google', 'Facebook']
CHAMBER_LABELS = {'LEG': 'State House', 'STSEN': 'State Senate'}
CHAMBER_ABBR = {'LEG': 'HD', 'STSEN': 'SD'}

DELTA_RED = 'FFDE5E4E'
DELTA_WHITE = 'FFFFFFFF'
DELTA_BLUE = 'FF3D6A91'

NAVY = '323b51'
WHITE = 'FFFFFF'
R_LIGHT, R_TOTAL, R_PARTY = 'fcefed', 'f2c1bb', 'de5e4e'
D_LIGHT, D_TOTAL, D_PARTY = 'ecf0f4', 'b4c6d5', '3d6a91'
NP_LIGHT, NP_TOTAL, NP_PARTY = 'fef6ed', 'fcddb8', 'd37518'
FOOTER_GREY = '6B7280'
FONT_TITLE = 'Superior Title'
FONT_BODY = 'Figtree'

PARTY_COLORS = {
    'R': (R_LIGHT, R_TOTAL, R_PARTY, 'REPUBLICAN PARTY TOTAL'),
    'D': (D_LIGHT, D_TOTAL, D_PARTY, 'DEMOCRAT PARTY TOTAL'),
    'NP': (NP_LIGHT, NP_TOTAL, NP_PARTY, 'NON-PARTISAN PARTY TOTAL'),
}
PARTY_ORDER = ['R', 'D', 'NP']

MEDIUM_BOTTOM = Border(bottom=Side(style='medium'))
DOUBLE_BOTTOM = Border(bottom=Side(style='double'))

GRID_SIDE = Side(style='thin', color='D9D9D9')
GRID = Border(left=GRID_SIDE, right=GRID_SIDE, top=GRID_SIDE, bottom=GRID_SIDE)
GRID_MEDIUM_BOTTOM = Border(left=GRID_SIDE, right=GRID_SIDE, top=GRID_SIDE, bottom=Side(style='medium'))
GRID_DOUBLE_BOTTOM = Border(left=GRID_SIDE, right=GRID_SIDE, top=GRID_SIDE, bottom=Side(style='double'))

WEEK_COL_RE = re.compile(r'^(\d{2}/\d{2})-(\d{2}/\d{2}) (DEM|GOP|total)$')


def fill(hexcolor):
    return PatternFill('solid', fgColor=hexcolor)


def font(bold=False, size=9, color=NAVY, name=FONT_BODY):
    return Font(name=name, size=size, bold=bold, color=color)


def _week_labels_from_header(columns):
    labels = []
    for col in columns:
        m = WEEK_COL_RE.match(col)
        if m and m.group(3) == 'DEM' and col not in labels:
            label = f'{m.group(1)}-{m.group(2)}'
            if label not in labels:
                labels.append(label)
    return labels


def _infer_years(week_labels, reference=None):
    """AdHawk week labels have no year. Anchor the last (most recent) week
    to today, then walk backward, decrementing the year every time we cross
    a Dec -> Jan boundary."""
    reference = reference or date.today()
    parsed = [(int(lbl[:2]), int(lbl[3:5])) for lbl in week_labels]  # (month, day) of week start
    years = [None] * len(parsed)

    last_m, last_d = parsed[-1]
    year = reference.year
    try:
        candidate = date(year, last_m, last_d)
    except ValueError:
        candidate = date(year, last_m, min(last_d, 28))
    if candidate > reference + timedelta(days=45):
        year -= 1
    elif candidate < reference - timedelta(days=400):
        year += 1
    years[-1] = year

    for i in range(len(parsed) - 2, -1, -1):
        m_this, _ = parsed[i]
        m_next, _ = parsed[i + 1]
        years[i] = years[i + 1] - 1 if m_this > m_next else years[i + 1]

    return years


def _natural_sort_key(code):
    parts = re.split(r'(\d+)', code)
    return tuple(int(p) if p.isdigit() else p for p in parts)


def _friendly_race_label(race, show_state_prefix):
    """'PA-LEG-13' -> 'State House District 13' (or 'PA State House District 13'
    when multiple states are present in one file, to avoid ambiguity). Falls
    back to the raw code for anything that doesn't match the LEG/STSEN pattern
    — e.g. a statewide race — since we don't know how to phrase those."""
    m = re.match(r'^([A-Za-z]{2,})-(LEG|STSEN)-0*(\d+)$', race)
    if not m:
        return race
    state, chamber, number = m.groups()
    label = f'{CHAMBER_LABELS[chamber]} District {number}'
    return f'{state} {label}' if show_state_prefix else label


def _tab_name(race, show_state_prefix):
    """'PA-LEG-13' -> 'HD-13' (or 'PA-HD-13' when multiple states are present).
    Falls back to the raw code, truncated to Excel's 31-char sheet-name limit,
    for anything that doesn't match the LEG/STSEN pattern."""
    m = re.match(r'^([A-Za-z]{2,})-(LEG|STSEN)-0*(\d+)$', race)
    if not m:
        return race[:31]
    state, chamber, number = m.groups()
    name = f'{CHAMBER_ABBR[chamber]}-{number}'
    return (f'{state}-{name}' if show_state_prefix else name)[:31]


def build_workbook(csv_path, out_path, reference_date=None, exclude_election_names=None):
    """exclude_election_names: optional list of Election Name values to drop before
    building tabs — for manually filtering out a known-bad/mistagged row. Not
    detected automatically; the caller has to have spotted it first."""
    df = pd.read_csv(csv_path)
    df = df.dropna(subset=['Election Name'])
    if exclude_election_names:
        df = df[~df['Election Name'].isin(exclude_election_names)]
    if df.empty:
        raise ValueError('No races found in this export (no rows with an Election Name).')

    week_labels = _week_labels_from_header(df.columns)
    if not week_labels:
        raise ValueError('Could not find any weekly DEM/GOP/total columns in this CSV.')
    years = _infer_years(week_labels, reference_date)
    week_headers = [f'{lbl.split("-")[0]}/{yr}' for lbl, yr in zip(week_labels, years)]

    races = sorted(df['Election Name'].unique(), key=_natural_sort_key)

    records_by_race = {}
    for race in races:
        sub = df[df['Election Name'] == race]
        recs = []
        for _, row in sub.iterrows():
            dem, gop = float(row.get('DEM', 0) or 0), float(row.get('GOP', 0) or 0)
            party = 'D' if dem > 0 else ('R' if gop > 0 else 'NP')
            party_col = 'DEM' if party == 'D' else ('GOP' if party == 'R' else 'total')
            weekly = []
            for wk in week_labels:
                col = f'{wk} {party_col}'
                val = row.get(col)
                weekly.append(float(val) if pd.notna(val) else 0.0)
            recs.append({
                'advertiser': row['Advertiser'],
                'party': party,
                'platform': row['Spend Platform'],
                'weekly': weekly,
            })
        records_by_race[race] = recs

    prefix_tokens = {r.split('-')[0] for r in races}
    if len(prefix_tokens) == 1:
        title_prefix = next(iter(prefix_tokens))
        report_title = f'{title_prefix} DIGITAL COMPETITIVE REPORT'
    else:
        report_title = 'DIGITAL COMPETITIVE REPORT'

    show_state_prefix = len(prefix_tokens) > 1
    sheet_names = {race: _tab_name(race, show_state_prefix) for race in races}

    last_col = 4 + len(week_labels)
    last_col_letter = get_column_letter(last_col)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def _col_pixel_width(width_units):
        return width_units * 7 + 5

    def _add_logo(ws, col_widths):
        if not LOGO_PATH.exists():
            return
        img = XLImage(str(LOGO_PATH))
        target_h = LOGO_HEIGHT_PX
        target_w = target_h * img.width / img.height

        remaining = target_w + LOGO_RIGHT_MARGIN_PX
        col_idx, col_off_px = 1, 0.0
        for i in range(len(col_widths), 0, -1):
            w_px = _col_pixel_width(col_widths[i - 1])
            if remaining <= w_px:
                col_idx, col_off_px = i, w_px - remaining
                break
            remaining -= w_px

        header_height_px = (34 + 18) * 4 / 3  # points -> pixels at 96dpi
        row_off_px = max((header_height_px - target_h) / 2, 0)

        img.width = target_w
        img.height = target_h
        marker = AnchorMarker(
            col=col_idx - 1, colOff=pixels_to_EMU(col_off_px),
            row=0, rowOff=pixels_to_EMU(row_off_px),
        )
        img.anchor = OneCellAnchor(
            _from=marker,
            ext=XDRPositiveSize2D(pixels_to_EMU(target_w), pixels_to_EMU(target_h)),
        )
        ws.add_image(img)

    def style_header_bar(ws, title_text, ncols, col_widths):
        # Merge the title cell across the FULL header width (not just a narrow
        # left block) — LibreOffice clips text at a merged cell's own boundary
        # regardless of whether neighboring cells are empty, so a narrow merge
        # truncates any title longer than that block, even though Excel itself
        # would have let it overflow into the empty cells beyond.
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=ncols)
        a1 = ws.cell(row=1, column=1, value=title_text)
        a1.font = font(bold=True, size=14, color=WHITE, name=FONT_TITLE)
        a1.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        for r in (1, 2):
            for col in range(1, ncols + 1):
                ws.cell(row=r, column=col).fill = fill(NAVY)
        ws.row_dimensions[1].height = 34
        ws.row_dimensions[2].height = 18
        _add_logo(ws, col_widths)

    def apply_print_setup(ws, ncols):
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = '1:3'

    race_col_widths = [34, 8, 16, 13] + [12] * len(week_labels)

    def set_col_widths(ws):
        for i, w in enumerate(race_col_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def build_race_sheet(race, recs):
        ws = wb.create_sheet(title=sheet_names[race])
        ncols = last_col
        style_header_bar(ws, sheet_names[race], ncols, race_col_widths)

        headers = ['CANDIDATE / COMMITTEE', 'PARTY', 'PLATFORM', 'TOTAL SPEND'] + week_headers
        for col, text in enumerate(headers, start=1):
            c = ws.cell(row=3, column=col, value=text)
            c.font = font(bold=True, size=9, color=WHITE)
            c.fill = fill(NAVY)
            c.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[3].height = 30

        set_col_widths(ws)
        ws.freeze_panes = 'E4'
        ws.sheet_view.showGridLines = False
        apply_print_setup(ws, ncols)

        row = 4
        party_label_row = {}

        for party in PARTY_ORDER:
            party_recs = [r for r in recs if r['party'] == party]
            if not party_recs:
                continue
            light, totalfill, party_fill, party_name = PARTY_COLORS[party]

            by_adv = {}
            for r in party_recs:
                by_adv.setdefault(r['advertiser'], []).append(r)
            adv_order = sorted(by_adv.keys(), key=lambda a: -sum(sum(x['weekly']) for x in by_adv[a]))
            adv_total_rows = []

            for adv in adv_order:
                plats = sorted(by_adv[adv], key=lambda r: PLATFORM_ORDER.get(r['platform'], 9))
                start_row = row
                for r in plats:
                    ws.cell(row=row, column=1, value=adv).font = font(bold=True, size=9)
                    ws.cell(row=row, column=1).fill = fill(light)
                    ws.cell(row=row, column=2, value=party).font = font(bold=True, size=9)
                    ws.cell(row=row, column=2).fill = fill(light)
                    ws.cell(row=row, column=2).alignment = Alignment(horizontal='center')
                    c3 = ws.cell(row=row, column=3, value=r['platform'])
                    c3.font = font(bold=False, size=9)
                    c3.fill = fill(light)
                    for i, wk_val in enumerate(r['weekly']):
                        col = 5 + i
                        cell = ws.cell(row=row, column=col, value=wk_val)
                        cell.font = font(bold=False, size=9)
                        cell.fill = fill(light)
                        cell.number_format = '$#,##0;-$#,##0;""'
                    dcell = ws.cell(row=row, column=4, value=f'=SUM(E{row}:{last_col_letter}{row})')
                    dcell.font = font(bold=False, size=9)
                    dcell.fill = fill(light)
                    dcell.number_format = '$#,##0;-$#,##0;""'
                    row += 1
                end_row = row - 1
                # Column A (advertiser name) merges across the block — purely
                # cosmetic, nothing reads it as data. Column B (party) stays
                # UNMERGED and gets a real value on every row, including the
                # total row below: the Summary tabs' SUMIFS formulas filter by
                # this column, and a merged cell only has a value in its
                # top-left anchor — every other row in the merge would read
                # back blank to a formula, silently dropping that platform's
                # spend from the party split (while a party-blind total still
                # matches, so the discrepancy allowed the bug to hide well).
                ws.merge_cells(start_row=start_row, start_column=1, end_row=row, end_column=1)

                tot_label = ws.cell(row=row, column=3, value=f'{adv} Total')
                tot_label.font = font(bold=True, size=9)
                tot_label.fill = fill(totalfill)
                tot_label.border = MEDIUM_BOTTOM
                ws.cell(row=row, column=1).fill = fill(totalfill)
                ws.cell(row=row, column=1).border = MEDIUM_BOTTOM
                pcell = ws.cell(row=row, column=2, value=party)
                pcell.font = font(bold=True, size=9)
                pcell.fill = fill(totalfill)
                pcell.alignment = Alignment(horizontal='center')
                pcell.border = MEDIUM_BOTTOM
                for i in range(len(week_labels)):
                    col = 5 + i
                    col_letter = get_column_letter(col)
                    cell = ws.cell(row=row, column=col, value=f'=SUM({col_letter}{start_row}:{col_letter}{end_row})')
                    cell.font = font(bold=True, size=9)
                    cell.fill = fill(totalfill)
                    cell.number_format = '$#,##0'
                    cell.border = MEDIUM_BOTTOM
                dcell = ws.cell(row=row, column=4, value=f'=SUM(E{row}:{last_col_letter}{row})')
                dcell.font = font(bold=True, size=9)
                dcell.fill = fill(totalfill)
                dcell.number_format = '$#,##0'
                dcell.border = MEDIUM_BOTTOM

                adv_total_rows.append(row)
                row += 2

            label_row = row
            party_label_row[party] = label_row
            ws.merge_cells(start_row=label_row, start_column=1, end_row=label_row, end_column=3)
            lab = ws.cell(row=label_row, column=1, value=party_name)
            lab.font = font(bold=True, size=10, color=WHITE)
            for col in range(1, 4):
                ws.cell(row=label_row, column=col).fill = fill(party_fill)
                ws.cell(row=label_row, column=col).border = MEDIUM_BOTTOM
            for i in range(len(week_labels)):
                col = 5 + i
                col_letter = get_column_letter(col)
                formula = '+'.join(f'{col_letter}{r}' for r in adv_total_rows)
                cell = ws.cell(row=label_row, column=col, value=f'={formula}')
                cell.font = font(bold=True, size=10, color=WHITE)
                cell.fill = fill(party_fill)
                cell.number_format = '$#,##0'
                cell.border = MEDIUM_BOTTOM
            dcell = ws.cell(row=label_row, column=4, value=f'=SUM(E{label_row}:{last_col_letter}{label_row})')
            dcell.font = font(bold=True, size=10, color=WHITE)
            dcell.fill = fill(party_fill)
            dcell.number_format = '$#,##0'
            dcell.border = MEDIUM_BOTTOM
            row = label_row + 2

        grand_row = row
        ws.merge_cells(start_row=grand_row, start_column=1, end_row=grand_row, end_column=3)
        lab = ws.cell(row=grand_row, column=1, value='GRAND TOTAL')
        lab.font = font(bold=True, size=10, color=WHITE)
        for col in range(1, 4):
            ws.cell(row=grand_row, column=col).fill = fill(NAVY)
            ws.cell(row=grand_row, column=col).border = DOUBLE_BOTTOM
        party_rows = list(party_label_row.values())
        for i in range(len(week_labels)):
            col = 5 + i
            col_letter = get_column_letter(col)
            formula = '+'.join(f'{col_letter}{r}' for r in party_rows)
            cell = ws.cell(row=grand_row, column=col, value=f'={formula}')
            cell.font = font(bold=True, size=10, color=WHITE)
            cell.fill = fill(NAVY)
            cell.number_format = '$#,##0'
            cell.border = DOUBLE_BOTTOM
        dcell = ws.cell(row=grand_row, column=4, value=f'=SUM(E{grand_row}:{last_col_letter}{grand_row})')
        dcell.font = font(bold=True, size=10, color=WHITE)
        dcell.fill = fill(NAVY)
        dcell.number_format = '$#,##0'
        dcell.border = DOUBLE_BOTTOM

        footer_row = grand_row + 2
        ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=last_col)
        f = ws.cell(row=footer_row, column=1, value='Report prepared by GPS Impact  |  Confidential')
        f.font = font(bold=False, size=8, color=FOOTER_GREY)

        return {'grand_row': grand_row, 'party_row': party_label_row}

    summary_col_widths = [28, 12, 14, 14, 14, 18]

    def build_summary_sheet(race_meta, sheet_name, tab_title, sum_col_letter, sheet_index):
        ws = wb.create_sheet(title=sheet_name, index=sheet_index)
        ncols = len(summary_col_widths)
        style_header_bar(ws, tab_title, ncols, summary_col_widths)

        headers = ['RACE', 'PLATFORM', 'TOTAL SPEND', 'GOP TOTAL', 'DEM TOTAL', 'DELTA (DEM-GOP)']
        for col, text in enumerate(headers, start=1):
            c = ws.cell(row=3, column=col, value=text)
            c.font = font(bold=True, size=9, color=WHITE)
            c.fill = fill(NAVY)
            c.alignment = Alignment(horizontal='center', vertical='center')
            c.border = GRID
        ws.row_dimensions[3].height = 24

        for i, w in enumerate(summary_col_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A4'
        ws.sheet_view.showGridLines = False
        apply_print_setup(ws, ncols)

        row = 4
        first_row = row
        subtotal_rows = []

        for race in races:
            meta = race_meta[race]
            grand_row = meta['grand_row']
            r_row = meta['party_row'].get('R')
            d_row = meta['party_row'].get('D')
            sheet_ref = sheet_names[race]
            label = _friendly_race_label(race, show_state_prefix)

            start_row = row
            for platform in SUMMARY_PLATFORMS:
                sum_range = f"'{sheet_ref}'!{sum_col_letter}$4:{sum_col_letter}${grand_row}"
                platform_range = f"'{sheet_ref}'!$C$4:$C${grand_row}"
                party_range = f"'{sheet_ref}'!$B$4:$B${grand_row}"

                ws.cell(row=row, column=2, value=platform).font = font(size=9)
                gop_cell = ws.cell(
                    row=row, column=4,
                    value=f'=SUMIFS({sum_range},{platform_range},"{platform}",{party_range},"R")',
                )
                dem_cell = ws.cell(
                    row=row, column=5,
                    value=f'=SUMIFS({sum_range},{platform_range},"{platform}",{party_range},"D")',
                )
                # No party filter here, so a non-partisan spender on this platform
                # (if any) still counts toward the total even though it has no
                # GOP/DEM column of its own.
                tot_cell = ws.cell(
                    row=row, column=3,
                    value=f'=SUMIFS({sum_range},{platform_range},"{platform}")',
                )
                delta_cell = ws.cell(row=row, column=6, value=f'=E{row}-D{row}')
                for c in (gop_cell, dem_cell, tot_cell, delta_cell):
                    c.font = font(size=9)
                    c.number_format = '$#,##0;-$#,##0;""'
                for col in range(1, ncols + 1):
                    ws.cell(row=row, column=col).border = GRID
                row += 1

            subtotal_row = row
            gop_ref = f"'{sheet_ref}'!{sum_col_letter}{r_row}" if r_row else '0'
            dem_ref = f"'{sheet_ref}'!{sum_col_letter}{d_row}" if d_row else '0'
            tot_ref = f"'{sheet_ref}'!{sum_col_letter}{grand_row}"

            ws.cell(row=subtotal_row, column=2, value='SUBTOTAL')
            gop_cell = ws.cell(row=subtotal_row, column=4, value=f'={gop_ref}')
            dem_cell = ws.cell(row=subtotal_row, column=5, value=f'={dem_ref}')
            tot_cell = ws.cell(row=subtotal_row, column=3, value=f'={tot_ref}')
            delta_cell = ws.cell(row=subtotal_row, column=6, value=f'=E{subtotal_row}-D{subtotal_row}')
            for col in range(1, ncols + 1):
                ws.cell(row=subtotal_row, column=col).fill = fill(D_TOTAL)
                ws.cell(row=subtotal_row, column=col).border = GRID_MEDIUM_BOTTOM
            for c in (gop_cell, dem_cell, tot_cell, delta_cell):
                c.font = font(bold=True, size=9)
                c.number_format = '$#,##0'
            ws.cell(row=subtotal_row, column=2).font = font(bold=True, size=9)

            label_cell = ws.cell(row=start_row, column=1, value=label)
            label_cell.font = font(bold=True, size=9)
            label_cell.alignment = Alignment(vertical='center', wrap_text=True)
            label_cell.hyperlink = f"#'{sheet_ref}'!A1"
            label_cell.border = GRID
            ws.merge_cells(start_row=start_row, start_column=1, end_row=subtotal_row, end_column=1)

            subtotal_rows.append(subtotal_row)
            row = subtotal_row + 2  # subtotal row + one blank separator

        last_data_row = row - 2
        ws.conditional_formatting.add(
            f'F{first_row}:F{last_data_row}',
            ColorScaleRule(
                start_type='min', start_color=DELTA_RED,
                mid_type='num', mid_value=0, mid_color=DELTA_WHITE,
                end_type='max', end_color=DELTA_BLUE,
            ),
        )

        total_row = row
        ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
        lab = ws.cell(row=total_row, column=1, value='TOTAL — ALL RACES')
        lab.font = font(bold=True, size=10, color=WHITE)
        for col in range(1, ncols + 1):
            ws.cell(row=total_row, column=col).fill = fill(NAVY)
            ws.cell(row=total_row, column=col).border = GRID_DOUBLE_BOTTOM
        for col in (3, 4, 5):
            col_letter = get_column_letter(col)
            formula = '+'.join(f'{col_letter}{r}' for r in subtotal_rows)
            cell = ws.cell(row=total_row, column=col, value=f'={formula}')
            cell.font = font(bold=True, size=10, color=WHITE)
            cell.number_format = '$#,##0'
        delta_total = ws.cell(row=total_row, column=6, value=f'=E{total_row}-D{total_row}')
        delta_total.font = font(bold=True, size=10, color=WHITE)
        delta_total.number_format = '$#,##0'

        footer_row = total_row + 2
        ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=ncols)
        f = ws.cell(row=footer_row, column=1, value='Report prepared by GPS Impact  |  Confidential')
        f.font = font(bold=False, size=8, color=FOOTER_GREY)

    race_meta = {race: build_race_sheet(race, records_by_race[race]) for race in races}
    build_summary_sheet(race_meta, 'Summary', report_title, 'D', 0)

    wb.save(out_path)
    return {'races': races, 'out_path': out_path}
